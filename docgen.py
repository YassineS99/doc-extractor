from docx import Document

from extractor import extract_spec
from pdf_reader import read_pdf_text

NOT_FOUND = "Not found in document"


def generate_docx(spec, output_path):
    """Build a Word document from a validated EquipmentSpec.
    Returns the list of fields that were missing."""
    doc = Document()
    doc.add_heading("Equipment Specification Sheet", level=1)
    doc.add_paragraph(f"Equipment: {spec.equipment_name}")

    rows = [
        ("Manufacturer", spec.manufacturer, ""),
        ("Model number", spec.model_number, ""),
        ("Rated power", spec.power_kw, "kW"),
        ("Supply voltage", spec.voltage_v, "V"),
        ("Net weight", spec.weight_kg, "kg"),
        ("Max operating temperature", spec.max_temp_c, "deg C"),
    ]

    table = doc.add_table(rows=1, cols=2)
    table.style = "Light Grid Accent 1"
    table.rows[0].cells[0].text = "Field"
    table.rows[0].cells[1].text = "Value"

    missing = []
    for label, value, unit in rows:
        cells = table.add_row().cells
        cells[0].text = label
        if value is None:
            cells[1].text = NOT_FOUND
            missing.append(label)
        else:
            cells[1].text = f"{value} {unit}".strip()

    doc.add_heading("Applicable Standards", level=2)
    if spec.standards:
        for standard in spec.standards:
            doc.add_paragraph(standard, style="List Bullet")
    else:
        doc.add_paragraph(NOT_FOUND)

    if missing:
        doc.add_paragraph("")
        doc.add_paragraph("Review needed. Missing fields: " + ", ".join(missing))

    doc.save(output_path)
    return missing


if __name__ == "__main__":
    text = read_pdf_text("sample_datasheet.pdf")
    spec = extract_spec(text)
    missing = generate_docx(spec, "output_spec.docx")
    print("Created output_spec.docx")
    print("Missing fields:", missing if missing else "none")